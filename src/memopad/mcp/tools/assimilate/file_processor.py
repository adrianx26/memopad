"""File content extraction utilities for various formats."""

import io
from typing import Optional

from loguru import logger

try:
    import pypdf
except ImportError:
    pypdf = None  # type: ignore

try:
    import docx
except ImportError:
    docx = None  # type: ignore

try:
    import openpyxl
except ImportError:
    openpyxl = None  # type: ignore

try:
    from PIL import Image
except ImportError:
    Image = None  # type: ignore


class FileProcessor:
    """Helper to extract text/metadata from various file formats."""

    @staticmethod
    def extract_pdf_text(content: bytes) -> str:
        """Extract text from PDF content."""
        if not pypdf:
            return "Error: pypdf not installed."
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            text = []
            for page in reader.pages:
                text.append(page.extract_text() or "")
            return "\n".join(text)
        except Exception as e:
            logger.error(f"Error extracting PDF: {e}")
            return f"Error extracting PDF: {e}"

    @staticmethod
    def extract_docx_text(content: bytes) -> str:
        """Extract text from DOCX content."""
        if not docx:
            return "Error: python-docx not installed."
        try:
            document = docx.Document(io.BytesIO(content))
            return "\n".join([para.text for para in document.paragraphs])
        except Exception as e:
            logger.error(f"Error extracting DOCX: {e}")
            return f"Error extracting DOCX: {e}"

    @staticmethod
    def extract_xlsx_text(content: bytes) -> str:
        """Extract text from XLSX content."""
        if not openpyxl:
            return "Error: openpyxl not installed."
        try:
            wb = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            )
            text = []
            for sheet in wb.worksheets:
                text.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = [str(cell) for cell in row if cell is not None]
                    if row_text:
                        text.append(" | ".join(row_text))
                text.append("\n")
            return "\n".join(text)
        except Exception as e:
            logger.error(f"Error extracting XLSX: {e}")
            return f"Error extracting XLSX: {e}"

    @staticmethod
    def extract_image_info(content: bytes) -> str:
        """Extract metadata from image content."""
        if not Image:
            return "Error: PIL not installed."
        try:
            img = Image.open(io.BytesIO(content))
            info = [
                f"Format: {img.format}",
                f"Size: {img.size} (Width x Height)",
                f"Mode: {img.mode}",
                f"Info: {img.info}",
            ]
            return "\n".join(info)
        except Exception as e:
            logger.error(f"Error processing image: {e}")
            return f"Error processing image: {e}"

    @classmethod
    def extract_text_content(
        cls, content: bytes, content_type: str, url: str
    ) -> str:
        """Dispatch to appropriate extractor based on content type/extension."""
        url_lower = url.lower()

        # PDF
        if "application/pdf" in content_type or url_lower.endswith(".pdf"):
            return cls.extract_pdf_text(content)

        # Word
        if (
            "wordprocessingml" in content_type
            or url_lower.endswith(".docx")
            or url_lower.endswith(".doc")
        ):
            if url_lower.endswith(".doc"):
                return "Error: .doc format not supported directly (only .docx)"
            return cls.extract_docx_text(content)

        # Excel
        if (
            "spreadsheet" in content_type
            or "excel" in content_type
            or url_lower.endswith(".xlsx")
            or url_lower.endswith(".xls")
        ):
            if url_lower.endswith(".xls"):
                return "Error: .xls format not supported directly (only .xlsx)"
            return cls.extract_xlsx_text(content)

        # CSV
        if "text/csv" in content_type or url_lower.endswith(".csv"):
            try:
                return content.decode("utf-8", errors="replace")
            except Exception:
                return "Error reading CSV"

        # Image
        if "image/" in content_type or url_lower.endswith(
            (".jpg", ".jpeg", ".png", ".bmp", ".gif", ".webp")
        ):
            return cls.extract_image_info(content)

        # Text/Code - try to decode as text
        try:
            return content.decode("utf-8", errors="replace")
        except Exception:
            return "Binary content (decoding failed)"
